import io
import json
import logging
import os
import zipfile

from docx import Document as Docx

from app.llm.base import LLMError
from tests.fakes import FakeProvider, SAMPLE_JD, SAMPLE_PROFILE

FAKE_KEY = "sk-ant-api03-SECRETSECRETSECRET1234567890"
H = {"X-LLM-Key": FAKE_KEY}


def make_docx() -> bytes:
    d = Docx()
    d.add_paragraph("Ada Lovelace - Software Engineer at Analytical Co 2019-Present")
    d.add_paragraph("Built REST APIs in Python.")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def sse_events(text: str):
    out = []
    for block in text.strip().split("\n\n"):
        ev, data = block.split("\n", 1)
        out.append((ev[7:], json.loads(data[6:])))
    return out


def test_health(client):
    assert client.get("/health").json() == {"status": "ok"}


def test_profile_upload_docx(client):
    r = client.post("/api/profile/parse", headers=H, files={"file": ("cv.docx", make_docx())})
    assert r.status_code == 200 and r.json()["contact"]["name"] == "Ada Lovelace"


def test_profile_upload_rejects_garbage(client):
    r = client.post("/api/profile/parse", headers=H, files={"file": ("cv.txt", b"hello world" * 10)})
    assert r.status_code == 422


def test_extract_match(client):
    r = client.post("/api/jd/extract", headers=H, json={"text": "Senior Backend Engineer at Acme" * 3})
    assert r.json()["company"] == "Acme Robotics"
    m = client.post("/api/match", headers=H, json={"profile": SAMPLE_PROFILE.model_dump(), "jd": SAMPLE_JD.model_dump()}).json()
    assert "Python" in m["matched"] and "Kubernetes" in m["missing"]
    assert m["semantic_score"] == 80 and 0 < m["score"] < 100


def test_generate_sse_flags_fake_claim(client):
    body = {"kind": "resume", "profile": SAMPLE_PROFILE.model_dump(), "jd": SAMPLE_JD.model_dump()}
    r = client.post("/api/generate", headers=H, json=body)
    evs = sse_events(r.text)
    assert evs[0][0] == "token" and evs[-1][0] == "done"
    done = evs[-1][1]
    assert done["doc"]["name"] == "Ada Lovelace"
    assert any("Kubernetes" in f["unsupported"] for f in done["flags"])


def test_regenerate_bullet(client):
    r = client.post("/api/regenerate-bullet", headers=H, json={
        "profile": SAMPLE_PROFILE.model_dump(), "jd": SAMPLE_JD.model_dump(), "bullet": "x"})
    assert r.json()["bullet"].startswith("Built Python APIs") and r.json()["unsupported"] == []


def _doc():
    return {"kind": "resume", "name": "Ada Lovelace", "contact": ["ada@example.com"],
            "sections": [{"title": "Experience", "paragraphs": [], "items": [
                {"heading": "Engineer", "subheading": "Analytical Co", "dates": "2019 - Present",
                 "note": "", "bullets": ["Built APIs"]}]}]}


def test_export_docx_all_templates(client):
    for t in ("classic", "modern", "compact"):
        r = client.post("/api/export/docx", json={"doc": _doc(), "template": t})
        assert r.status_code == 200
        assert "Built APIs" in "\n".join(p.text for p in Docx(io.BytesIO(r.content)).paragraphs)
        assert zipfile.is_zipfile(io.BytesIO(r.content))


def test_export_pdf_all_templates(client):
    for t in ("classic", "modern", "compact"):
        r = client.post("/api/export/pdf", json={"doc": _doc(), "template": t})
        if r.status_code == 503:  # WeasyPrint system libs missing: graceful failure
            assert "unavailable" in r.json()["detail"]
            continue
        assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_pdf_graceful_when_unavailable(client, monkeypatch):
    from app import main
    from app.export.pdf import PdfUnavailable

    def boom(*a):
        raise PdfUnavailable("PDF export is unavailable: libs missing")

    monkeypatch.setattr(main, "export_pdf", boom)
    r = client.post("/api/export/pdf", json={"doc": _doc()})
    assert r.status_code == 503


def test_key_never_in_logs_or_responses(client, provider, caplog):
    caplog.set_level(logging.DEBUG)
    logging.getLogger("tailorurresume").warning("debug header X-LLM-Key: %s and key %s", FAKE_KEY, FAKE_KEY)
    provider.fail_with = LLMError(f"upstream said bad key {FAKE_KEY}")
    responses = [
        client.post("/api/test-connection", headers=H),
        client.post("/api/match", headers=H, json={"profile": SAMPLE_PROFILE.model_dump(), "jd": SAMPLE_JD.model_dump()}),
        client.post("/api/generate", headers=H, json={"kind": "cv", "profile": SAMPLE_PROFILE.model_dump(), "jd": SAMPLE_JD.model_dump()}),
    ]
    provider.fail_with = RuntimeError(f"boom {FAKE_KEY}")
    responses.append(client.post("/api/match", headers=H, json={"profile": SAMPLE_PROFILE.model_dump(), "jd": SAMPLE_JD.model_dump()}))
    provider.fail_with = None
    responses.append(client.post("/api/test-connection", headers=H))
    for r in responses:
        assert "SECRETSECRET" not in r.text
        assert "SECRETSECRET" not in json.dumps(dict(r.headers))
    assert responses[-1].json() == {"ok": True, "provider": "fake", "model": "fake-model"}
    assert caplog.records
    for rec in caplog.records:
        assert "SECRETSECRET" not in rec.getMessage()
        assert "SECRETSECRET" not in (rec.exc_text or "")
    assert "SECRETSECRET" not in caplog.text


def test_no_files_written_in_full_run(client, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("TMPDIR", str(tmp_path))
    import tempfile
    tempfile.tempdir = str(tmp_path)
    try:
        before = set(os.listdir(tmp_path))
        client.post("/api/profile/parse", headers=H, files={"file": ("cv.docx", make_docx())})
        client.post("/api/jd/extract", headers=H, json={"text": "Senior Backend Engineer at Acme" * 3})
        client.post("/api/match", headers=H, json={"profile": SAMPLE_PROFILE.model_dump(), "jd": SAMPLE_JD.model_dump()})
        client.post("/api/generate", headers=H, json={"kind": "resume", "profile": SAMPLE_PROFILE.model_dump(), "jd": SAMPLE_JD.model_dump()})
        client.post("/api/export/docx", json={"doc": _doc()})
        client.post("/api/export/pdf", json={"doc": _doc()})
        assert set(os.listdir(tmp_path)) == before
    finally:
        tempfile.tempdir = None
