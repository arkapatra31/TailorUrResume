from __future__ import annotations

import io

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from ..schemas import Document
from .styles import TEMPLATES


def _rule(paragraph, color: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color)):
        bottom.set(qn(k), v)
    b.append(bottom)
    pPr.append(b)


def export_docx(doc: Document, template: str) -> io.BytesIO:
    t = TEMPLATES[template]
    d = DocxDocument()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Mm(float(t["margin"].replace("mm", ""))))
    normal = d.styles["Normal"]
    normal.font.name = t["docx_font"]
    normal.font.size = Pt(float(t["size"].replace("pt", "")) + 0.4)
    normal.paragraph_format.space_after = Pt(2)
    accent = RGBColor(*t["docx_accent"])
    hexcolor = "%02X%02X%02X" % t["docx_accent"]
    align = WD_ALIGN_PARAGRAPH.CENTER if t["docx_center"] else WD_ALIGN_PARAGRAPH.LEFT
    width = sec.page_width - sec.left_margin - sec.right_margin

    if doc.name:
        p = d.add_paragraph()
        p.alignment = align
        r = p.add_run(doc.name)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(float(t["h1"].replace("pt", ""))), accent
    if doc.contact:
        p = d.add_paragraph("  |  ".join(doc.contact))
        p.alignment = align
        p.paragraph_format.space_after = Pt(6)

    for s in doc.sections:
        if s.title and doc.kind != "cover_letter":
            p = d.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            r = p.add_run(s.title.upper())
            r.bold, r.font.color.rgb = True, accent
            _rule(p, hexcolor)
        for para in s.paragraphs:
            p = d.add_paragraph(para)
            if doc.kind == "cover_letter":
                p.paragraph_format.space_after = Pt(9)
        for it in s.items:
            if it.heading or it.dates:
                p = d.add_paragraph()
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
                p.add_run(it.heading).bold = True
                if it.subheading:
                    p.add_run(" — " + it.subheading).italic = True
                if it.dates:
                    p.add_run("\t" + it.dates)
            if it.note:
                d.add_paragraph(it.note)
            for b in it.bullets:
                bp = d.add_paragraph(b, style="List Bullet")
                bp.paragraph_format.space_after = Pt(1)
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)
    return buf
