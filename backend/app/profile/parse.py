"""Resume upload parsing. Everything happens in memory."""
from __future__ import annotations

import io

from docx import Document as DocxDocument
from pypdf import PdfReader

from .. import config
from ..llm.base import LLMProvider
from ..schemas import Profile


class ParseError(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if len(data) > config.MAX_UPLOAD_BYTES:
        raise ParseError("File is too large.")
    try:
        if name.endswith(".pdf") or data[:5] == b"%PDF-":
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ParseError("Encrypted PDFs are not supported.")
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        elif name.endswith(".docx") or data[:2] == b"PK":
            doc = DocxDocument(io.BytesIO(data))
            lines = [p.text for p in doc.paragraphs]
            for t in doc.tables:
                for row in t.rows:
                    lines.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(lines)
        else:
            raise ParseError("Unsupported file type. Upload a PDF or DOCX.")
    except ParseError:
        raise
    except Exception:  # noqa: BLE001
        raise ParseError("Could not read that file. Is it a valid PDF/DOCX?") from None
    text = text.strip()
    if len(text) < 30:
        raise ParseError("No text found in the file (scanned image PDFs are not supported).")
    return text[: config.MAX_TEXT_CHARS]


SYSTEM = (
    "You parse resumes into structured JSON. Copy facts exactly as written; never invent or "
    "embellish. Keep every bullet from the source. Dates as written (e.g. 'Jan 2021', 'Present'). "
    "Put all skills as a flat list of short strings. Leave unknown fields empty."
)


async def parse_profile(provider: LLMProvider, text: str) -> Profile:
    profile = await provider.complete_json(SYSTEM, f"Resume text:\n\n{text}", Profile)
    profile.attested_skills = []  # only the user can approve skill bridges, never the parser
    return profile
